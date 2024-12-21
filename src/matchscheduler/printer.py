import itertools
from datetime import datetime
from pathlib import Path

from icalendar import Calendar, Event
from openpyxl import Workbook

from .match import Match
from .season import Season


class Printer:

    def __init__(self, season: Season):
        self.season = season

    # export schedule into a excel file with each round as a row
    # consisting of the date in the first column and in every next column a match
    def export(self, folderpath: Path) -> None:
        self.export_excel(folderpath)
        self.export_calendar(folderpath)

    def export_excel(self, folderpath: Path) -> None:
        all_dates = sorted(
            [r.day for r in self.season.schedule.rounds] + self.season.excluded_dates
        )
        excel = Workbook()
        sheet = excel.active
        sheet.title = "Schedule"  # type: ignore
        sheet.append(  # type: ignore
            ["Date"] + [f"Match {i+1}" for i in range(self.season.num_courts)]
        )
        for d in all_dates:
            if d in self.season.excluded_dates:
                sheet.append([str(d)])  # type: ignore
            else:
                round = next(filter(lambda r: r.day == d, self.season.schedule.rounds))
                sheet.append(  # type: ignore
                    [str(round.day)] + [str(m) for m in round.matches]
                )  # type: ignore

        # add an additional sheet to excel workbook with columns for each player
        # and their match partners
        sheet = excel.create_sheet("Partner by Player")  # type: ignore
        sheet.append(["Date"] + [str(p) for p in self.season.players])  # type: ignore
        for round in self.season.schedule.rounds:
            row = [str(round.day)]
            matches = round.matches
            for player in self.season.players:
                append_string = ""
                for m in matches:
                    if player in m.get_players():
                        opponent = m.player1 if m.player1 != player else m.player2
                        append_string += str(opponent.name if opponent is not None else "...")
                        break
                row.append(append_string)  # type: ignore
            sheet.append(row)  # type: ignore

        # add an additional sheet to excel workbook with columns for each possible match
        # and each row marks with an x if the match is played on that day
        sheet = excel.create_sheet("Matches Overview")  # type: ignore
        possible_matches = list(
            Match(p, q) for p, q in itertools.combinations(self.season.players, 2)
        )
        sheet.append(["Date"] + [str(m) for m in possible_matches])  # type: ignore
        for round in self.season.schedule.rounds:
            row = [str(round.day)]
            for match in possible_matches:
                append_string = ""
                for m in round.matches:
                    if m == match:
                        append_string = "x"
                        break
                row.append(append_string)  # type: ignore
            sheet.append(row)  # type: ignore

        sheet = excel.create_sheet("Costs")
        sheet.append([""] + [str(p) for p in self.season.players])  # type: ignore
        cost_per_match = (
            self.season.overall_cost
            / (len(self.season.schedule.rounds) * self.season.num_courts)
            / 2
        )
        sheet.append(  # type: ignore
            ["Matches"]
            + [
                len(self.season.schedule.get_match_indizes_of_player(p))
                for p in self.season.players
            ]
        )
        sheet.append(  # type: ignore
            ["Cost"]
            + [
                len(self.season.schedule.get_match_indizes_of_player(p)) * cost_per_match
                for p in self.season.players
            ]
        )

        excel.save(folderpath / "schedule.xlsx")

    def export_calendar(self, folderpath: Path) -> None:
        # create a calendar for each player with his matches
        for i, p in enumerate(self.season.players):
            cal = Calendar()
            cal.add("prodid", "-//MatchScheduler//MatchScheduler//EN")
            cal.add("version", "2.0")
            cal.add("name", f"MatchScheduler - {p.name}")
            cal.add("X-WR-CALNAME", f"MatchScheduler - {p.name}")
            cal.add("X-WR-TIMEZONE", "Europe/Vienna")
            cal.add("X-WR-CALDESC", f"MatchScheduler - {p.name}")
            for round_index, match_index in self.season.schedule.get_match_indizes_of_player(p):
                round = self.season.schedule.rounds[round_index]
                match = round.matches[match_index]
                event = Event()
                event.add("summary", self.season.calendar_title)
                event.add("description", str(match))
                event.add(
                    "dtstart",
                    datetime.combine(round.day, self.season.time_start),
                )
                event.add("dtend", datetime.combine(round.day, self.season.time_end))
                cal.add_component(event)
            with open(folderpath / f"{p.name}.ics", "wb") as f:
                f.write(cal.to_ical())
